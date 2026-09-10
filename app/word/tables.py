from __future__ import annotations

from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from app.models import TableData


def _set_borders(table) -> None:
    properties = table._tbl.tblPr
    existing = properties.find(qn("w:tblBorders"))
    if existing is not None:
        properties.remove(existing)
    borders = OxmlElement("w:tblBorders")
    for name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = OxmlElement(f"w:{name}")
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "4")
        edge.set(qn("w:color"), "D9D9D9")
        borders.append(edge)
    properties.append(borders)


def add_table(document, data: TableData):
    table = document.add_table(rows=data.rows, cols=data.cols)
    table.autofit = True
    _set_borders(table)
    for item in sorted(data.cells, key=lambda cell: (cell.row, cell.col)):
        cell = table.cell(item.row, item.col)
        if item.row_span > 1 or item.col_span > 1:
            cell = cell.merge(
                table.cell(item.row + item.row_span - 1, item.col + item.col_span - 1)
            )
        cell.text = item.text
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.space_after = Pt(0)
            for run in paragraph.runs:
                run.font.name = "Microsoft YaHei"
                run.font.size = Pt(9)
    return table
