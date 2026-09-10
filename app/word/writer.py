from __future__ import annotations

from pathlib import Path
from io import BytesIO
from typing import Iterable

from docx import Document
from docx.shared import Inches, Pt

from app.models import PageContent
from app.word.tables import add_table


class WordWriter:
    """Write normalized page content to an editable DOCX."""

    def write(self, pages: Iterable[PageContent], output_path: Path) -> Path:
        page_list = list(pages)
        document = Document()
        normal = document.styles["Normal"]
        normal.font.name = "Microsoft YaHei"
        normal.font.size = Pt(10.5)

        for page_index, page in enumerate(page_list):
            items = (
                [("text", block) for block in page.text_blocks]
                + [("table", table) for table in page.tables]
                + [("image", image) for image in page.images]
            )
            for kind, item in sorted(items, key=lambda pair: (pair[1].bbox[1], pair[1].bbox[0])):
                if kind == "text":
                    paragraph = document.add_paragraph(item.text)
                    paragraph.paragraph_format.space_after = Pt(2)
                elif kind == "table":
                    add_table(document, item)
                else:
                    paragraph = document.add_paragraph()
                    width_inches = min(
                        max((item.bbox[2] - item.bbox[0]) / 72.0, 0.5), 6.5
                    )
                    paragraph.add_run().add_picture(
                        BytesIO(item.data), width=Inches(width_inches)
                    )
            if page.error:
                document.add_paragraph(f"[第 {page.page_number} 页转换失败：{page.error}]")
            if page_index < len(page_list) - 1:
                document.add_page_break()

        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        document.save(output_path)
        return output_path
