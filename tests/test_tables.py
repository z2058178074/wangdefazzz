from pathlib import Path

import numpy as np
from docx import Document

from app.models import TableCell, TableData, TextBlock
from app.table.assign import assign_blocks
from app.table.grid import GridDetector
from app.table.raster import RasterGridDetector
from app.word.tables import add_table


def regular_grid_lines():
    return [
        (0, 0, 100, 0),
        (0, 30, 100, 30),
        (0, 60, 100, 60),
        (0, 0, 0, 60),
        (50, 0, 50, 60),
        (100, 0, 100, 60),
    ]


def test_grid_detector_recovers_rows_columns_and_cells() -> None:
    table = GridDetector.from_lines(regular_grid_lines(), (100, 60))
    assert table is not None
    assert (table.rows, table.cols) == (2, 2)
    assert [(cell.row, cell.col) for cell in table.cells] == [(0, 0), (0, 1), (1, 0), (1, 1)]


def test_missing_vertical_header_segment_creates_merged_cell() -> None:
    lines = [line for line in regular_grid_lines() if line != (50, 0, 50, 60)]
    lines.append((50, 30, 50, 60))

    table = GridDetector.from_lines(lines, (100, 60))

    header = next(cell for cell in table.cells if cell.row == 0 and cell.col == 0)
    assert header.col_span == 2
    assert not any(cell.row == 0 and cell.col == 1 for cell in table.cells)


def test_assigns_chinese_and_numbers_to_cells() -> None:
    table = GridDetector.from_lines(regular_grid_lines(), (100, 60))
    blocks = [
        TextBlock("项目", (5, 5, 35, 20)),
        TextBlock("金额", (55, 5, 85, 20)),
        TextBlock("营业收入", (5, 35, 45, 50)),
        TextBlock("1,234.56", (55, 35, 95, 50)),
    ]

    assigned = assign_blocks(table, blocks)

    assert [cell.text for cell in assigned.cells] == ["项目", "金额", "营业收入", "1,234.56"]


def test_raster_detector_finds_scanned_grid() -> None:
    import cv2

    image = np.full((240, 400, 3), 255, dtype=np.uint8)
    for x in (20, 200, 380):
        cv2.line(image, (x, 20), (x, 220), (0, 0, 0), 3)
    for y in (20, 120, 220):
        cv2.line(image, (20, y), (380, y), (0, 0, 0), 3)

    table = RasterGridDetector().detect(image)

    assert table is not None
    assert table.rows == 2
    assert table.cols == 2


def test_word_table_preserves_merge_and_borders(tmp_path: Path) -> None:
    data = TableData(
        rows=2,
        cols=2,
        bbox=(0, 0, 100, 60),
        cells=[
            TableCell(0, 0, (0, 0, 100, 30), "营业收入审定表", col_span=2),
            TableCell(1, 0, (0, 30, 50, 60), "本期"),
            TableCell(1, 1, (50, 30, 100, 60), "1,234.56"),
        ],
    )
    doc = Document()
    add_table(doc, data)
    output = tmp_path / "table.docx"
    doc.save(output)

    reopened = Document(output)
    assert len(reopened.tables) == 1
    assert reopened.tables[0].cell(0, 0).text == "营业收入审定表"
    assert reopened.tables[0].cell(0, 0)._tc is reopened.tables[0].cell(0, 1)._tc
    xml = reopened.tables[0]._tbl.xml
    assert "tblBorders" in xml
    assert "D9D9D9" in xml
