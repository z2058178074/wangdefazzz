from pathlib import Path

import cv2
import numpy as np
from PIL import Image
from docx import Document
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas

from app.converter import ConversionOptions, PDFConverter
from app.models import TextBlock


def make_financial_table(path: Path) -> None:
    pdf = canvas.Canvas(str(path), pagesize=(500, 300))
    xs = (40, 250, 460)
    ys = (250, 200, 150)
    for x in xs:
        pdf.line(x, ys[-1], x, ys[0])
    for y in ys:
        pdf.line(xs[0], y, xs[-1], y)
    pdf.drawString(55, 220, "Item")
    pdf.drawString(270, 220, "Amount")
    pdf.drawString(55, 170, "Revenue")
    pdf.drawString(270, 170, "1,234.56")
    pdf.save()


def test_text_pdf_table_becomes_editable_word_table(tmp_path: Path) -> None:
    source = tmp_path / "revenue-audit.pdf"
    make_financial_table(source)

    result = PDFConverter().convert_file(
        source,
        ConversionOptions(enable_ocr=False, preserve_tables=True),
    )

    assert result.success
    doc = Document(result.output_path)
    assert len(doc.tables) == 1
    values = [[cell.text for cell in row.cells] for row in doc.tables[0].rows]
    assert values == [["Item", "Amount"], ["Revenue", "1,234.56"]]


def test_scanned_financial_grid_becomes_editable_word_table(tmp_path: Path) -> None:
    source = tmp_path / "scanned-financial-table.pdf"
    pixels = np.full((240, 400, 3), 255, dtype=np.uint8)
    for x in (20, 200, 380):
        cv2.line(pixels, (x, 20), (x, 220), (0, 0, 0), 3)
    for y in (20, 120, 220):
        cv2.line(pixels, (20, y), (380, y), (0, 0, 0), 3)
    pdf = canvas.Canvas(str(source), pagesize=(400, 240))
    pdf.drawImage(ImageReader(Image.fromarray(pixels)), 0, 0, width=400, height=240)
    pdf.save()

    class FakeOCR:
        def recognize(self, image):
            height, width = image.shape[:2]
            sx, sy = width / 400, height / 240
            return [
                TextBlock("项目", (50 * sx, 50 * sy, 100 * sx, 80 * sy)),
                TextBlock("金额", (230 * sx, 50 * sy, 280 * sx, 80 * sy)),
                TextBlock("营业收入", (50 * sx, 150 * sy, 140 * sx, 180 * sy)),
                TextBlock("9,876.54", (230 * sx, 150 * sy, 330 * sx, 180 * sy)),
            ]

    result = PDFConverter(ocr_engine=FakeOCR()).convert_file(source)

    assert result.success
    doc = Document(result.output_path)
    assert len(doc.tables) == 1
    assert [[cell.text for cell in row.cells] for row in doc.tables[0].rows] == [
        ["项目", "金额"],
        ["营业收入", "9,876.54"],
    ]
